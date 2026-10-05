import io

import pytest
from docx import Document
from fpdf import FPDF

from app.builder.export_pdf import render_pdf
from app.parsers import DocumentTooLarge, UnsupportedDocument, parse_upload
from app.parsers.docx_parser import parse_docx
from app.parsers.pdf_parser import parse_pdf
from app.schemas.resume import ResumeDocument

LONG_BULLET = (
    "Partnered with finance and operations leaders to redesign the monthly forecasting workflow, "
    "improving accuracy by 18% and cutting close time by three days."
)


def _pdf_pages(build) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    build(pdf)
    return bytes(pdf.output())


def test_dense_single_column_page_is_not_flagged_multi_column():
    # Regression: the original midline word count flagged exactly this layout.
    doc = ResumeDocument.model_validate({
        "basics": {"name": "A B", "email": "a@b.co"},
        "experience": [
            {"title": "Analyst", "company": f"Co{i}", "start": "2019", "end": "2020", "bullets": [LONG_BULLET] * 5}
            for i in range(4)
        ],
    })
    parsed = parse_pdf(render_pdf(doc))
    assert parsed["page_count"] == 2
    assert parsed["multi_column"] is False


@pytest.mark.parametrize("sidebar_mm", [45, 60, 90])
def test_two_column_layout_is_flagged(sidebar_mm):
    def build(pdf):
        for i in range(30):
            pdf.set_xy(10, 20 + i * 8)
            pdf.cell(sidebar_mm, 6, f"Skill item {i} SQL Python")
            pdf.set_xy(10 + sidebar_mm + 8, 20 + i * 8)
            pdf.cell(190 - sidebar_mm - 8, 6, f"Led the analytics team for project {i} delivering results")

    assert parse_pdf(_pdf_pages(build))["multi_column"] is True


def test_right_aligned_dates_are_not_flagged_multi_column():
    def build(pdf):
        for i in range(6):
            pdf.cell(130, 6, f"Data Analyst, Company {i}")
            pdf.cell(60, 6, "Jan 2020 - Present", align="R", new_x="LMARGIN", new_y="NEXT")
            for _ in range(3):
                pdf.multi_cell(190, 6, "- " + LONG_BULLET, new_x="LMARGIN", new_y="NEXT")

    assert parse_pdf(_pdf_pages(build))["multi_column"] is False


def test_page_limit_is_enforced_before_parsing_pages():
    def build(pdf):
        for _ in range(4):
            pdf.add_page()
            pdf.cell(0, 10, "page")

    with pytest.raises(DocumentTooLarge, match="5 pages"):
        parse_pdf(_pdf_pages(build), max_pages=3)


def test_docx_list_bullets_are_recovered_as_bullets():
    # Regression: Word bullets are numbering, not text, and were scored as plain lines.
    document = Document()
    document.add_paragraph("Experience")
    document.add_paragraph("Led migration to Power BI, saving 10 hours per week.", style="List Bullet")
    buffer = io.BytesIO()
    document.save(buffer)

    text = parse_docx(buffer.getvalue())["text"]
    assert "• Led migration to Power BI" in text
    assert "• Experience" not in text


def test_parse_upload_checks_magic_bytes_not_just_extension():
    with pytest.raises(UnsupportedDocument):
        parse_upload("resume.pdf", b"PK\x03\x04 actually a zip")
    with pytest.raises(UnsupportedDocument):
        parse_upload("resume.docx", b"%PDF-1.7 actually a pdf")
    with pytest.raises(UnsupportedDocument):
        parse_upload("resume.txt", b"plain text")
