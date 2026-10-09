"""Single-column DOCX export with Word's own list bullets and no tables,
text boxes, headers or footers: the structures ATS parsers handle worst.

Template styling uses only paragraph-level formatting a parser ignores
safely: right-aligned dates sit behind a right tab stop (not in a table
cell), heading rules are paragraph borders (not drawn shapes). Word can't
embed the bundled fonts, so each font maps to a close face every Office
install has (see "docx" in templates.json).
"""

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_UNDERLINE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.builder.export_pdf import HEADING_SCALE, NAME_LEADING, SMALLCAPS_SCALE
from app.builder.layout import build_blocks
from app.builder.templates import FONTS, resolve_style
from app.schemas.resume import ResumeDocument

TEXT = RGBColor(0x1F, 0x29, 0x37)
STRONG = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x4B, 0x55, 0x63)


def _bottom_border(paragraph, rgb: tuple[int, int, int], size_eighths: int, kind: str = "single") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), kind)
    bottom.set(qn("w:sz"), str(size_eighths))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "{:02X}{:02X}{:02X}".format(*rgb))
    borders.append(bottom)
    p_pr.append(borders)


def render_docx(doc: ResumeDocument) -> bytes:
    style = resolve_style(doc)
    size = style.font_size
    page_w, page_h = style.page_mm
    text_width = Mm(page_w - 2 * style.margin)
    accent = RGBColor(*style.accent_rgb)
    header_align = WD_ALIGN_PARAGRAPH.CENTER if style.header_align == "center" else WD_ALIGN_PARAGRAPH.LEFT

    document = Document()
    for section in document.sections:
        section.page_width, section.page_height = Mm(page_w), Mm(page_h)
        section.top_margin = section.bottom_margin = Mm(style.margin)
        section.left_margin = section.right_margin = Mm(style.margin)

    font_name = FONTS[style.font]["docx"]
    normal = document.styles["Normal"]
    normal.font.name = font_name
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    normal.font.size = Pt(size)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    def paragraph(*, style_name: str | None = None, leading_pt: float | None = None,
                  space_before: float = 0, align=None):
        para = document.add_paragraph(style=style_name)
        fmt = para.paragraph_format
        fmt.space_before = Pt(space_before)
        fmt.space_after = Pt(0)
        fmt.line_spacing = Pt(leading_pt or size * style.line_height)
        fmt.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        if align is not None:
            fmt.alignment = align
        return para

    def run(para, text: str, *, pt: float = size, bold=False, italic=False, color: RGBColor = TEXT):
        r = para.add_run(text)
        r.bold, r.italic = bold, italic
        r.font.size = Pt(pt)
        r.font.color.rgb = color
        return r

    previous = None
    for block in build_blocks(doc, style):
        kind = block.kind
        if kind == "name":
            para = paragraph(leading_pt=style.name_size * NAME_LEADING, align=header_align)
            run(para, block.text, pt=style.name_size, bold=True, color=accent)
        elif kind == "headline":
            run(paragraph(align=header_align), block.text, pt=size + 1.5,
                color=accent if style.headline_color == "accent" else TEXT)
        elif kind == "contact":
            run(paragraph(align=header_align), block.text, pt=size - 0.5, color=MUTED)
        elif kind == "heading":
            heading_pt = size * HEADING_SCALE
            para = paragraph(
                leading_pt=heading_pt * 1.25,
                space_before=size * style.line_height * 0.75 + 2,
                align=WD_ALIGN_PARAGRAPH.CENTER if style.heading_align == "center" else None,
            )
            para.paragraph_format.space_after = Pt(3)
            para.paragraph_format.keep_with_next = True
            runs = []
            if style.heading_case == "smallcaps":
                for i, word in enumerate(block.text.split(" ")):
                    if i:
                        runs.append(run(para, " ", pt=heading_pt * SMALLCAPS_SCALE, bold=True, color=accent))
                    runs.append(run(para, word[:1], pt=heading_pt, bold=True, color=accent))
                    if word[1:]:
                        runs.append(run(para, word[1:], pt=heading_pt * SMALLCAPS_SCALE, bold=True, color=accent))
            else:
                runs.append(run(para, block.text, pt=heading_pt, bold=True, color=accent))
            if style.heading_style == "short":
                # Word borders span the column, so the short bar becomes a thick underline.
                for r in runs:
                    r.font.underline = WD_UNDERLINE.THICK
            elif style.heading_style == "double":
                _bottom_border(para, style.rule_rgb, 6, kind="double")
            elif style.heading_style != "plain":
                _bottom_border(para, style.rule_rgb, 6 if style.heading_style == "rule" else 8)
        elif kind == "entry_title":
            gap = 0 if previous is not None and previous.kind == "heading" else 3
            para = paragraph(space_before=gap)
            para.paragraph_format.keep_with_next = True
            if block.aside:
                para.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
            run(para, block.text, bold=True, color=STRONG)
            if block.aside:
                run(para, "\t" + block.aside, italic=True, color=MUTED)
        elif kind == "entry_meta":
            run(paragraph(), block.text, pt=size - 0.5, italic=True, color=MUTED)
        elif kind == "bullet":
            run(paragraph(style_name="List Bullet"), block.text)
        elif kind == "skill":
            para = paragraph()
            run(para, f"{block.label}:", bold=True)
            run(para, block.text[len(block.label) + 1:])
        else:
            run(paragraph(), block.text)
        previous = block

    document.core_properties.title = doc.basics.name or "Resume"
    document.core_properties.author = doc.basics.name or ""
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
